"""Build docs/management-proof/learnwise-edvidura-deck.pptx — same story as the HTML deck."""
from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "management-proof" / "learnwise-edvidura-deck.pptx"

INK = RGBColor(0x03, 0x20, 0x19)
MUTED = RGBColor(0x3C, 0x4A, 0x42)
PRIMARY = RGBColor(0x00, 0x6C, 0x4B)
MINT = RGBColor(0x34, 0xD3, 0x99)
SKY = RGBColor(0x40, 0xC2, 0xFD)
CITRUS = RGBColor(0xF0, 0xC0, 0x1A)
CORAL = RGBColor(0xF0, 0x5A, 0x3A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SURFACE = RGBColor(0xE6, 0xFF, 0xF5)
WARM_BG = RGBColor(0xFF, 0xF5, 0xD6)
DEEP = RGBColor(0x0E, 0x2B, 0x24)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)


def _solid_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _set_run(run, text: str, *, size: int, bold: bool = False, color: RGBColor = INK) -> None:
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def add_text(slide, left, top, width, height, lines, *, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        text, size, bold, color = line
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        run = p.add_run()
        _set_run(run, text, size=size, bold=bold, color=color)
    return box


def add_rect(slide, left, top, width, height, fill: RGBColor, *, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1.5)
    try:
        shape.adjustments[0] = 0.15
    except Exception:
        pass
    return shape


def add_card_text(shape, text: str, *, title: str | None = None, size=14, color=INK, title_color=PRIMARY):
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if title:
        run = p.add_run()
        _set_run(run, title, size=size + 4, bold=True, color=title_color)
        p2 = tf.add_paragraph()
        p2.space_before = Pt(6)
        run2 = p2.add_run()
        _set_run(run2, text, size=size, bold=False, color=color)
    else:
        run = p.add_run()
        _set_run(run, text, size=size, bold=False, color=color)


def add_appear_animation(slide, shapes_in_order) -> None:
    if not shapes_in_order:
        return
    for child in list(slide._element):
        if child.tag == qn("p:timing"):
            slide._element.remove(child)

    timing = etree.SubElement(slide._element, qn("p:timing"))
    tnLst = etree.SubElement(timing, qn("p:tnLst"))
    par = etree.SubElement(tnLst, qn("p:par"))
    cTn = etree.SubElement(par, qn("p:cTn"), {"id": "1", "dur": "indefinite", "restart": "never", "nodeType": "tmRoot"})
    childTnLst = etree.SubElement(cTn, qn("p:childTnLst"))
    seq = etree.SubElement(childTnLst, qn("p:seq"), {"concurrent": "1", "nextAc": "seek"})
    seqCTn = etree.SubElement(seq, qn("p:cTn"), {"id": "2", "dur": "indefinite", "nodeType": "mainSeq"})
    seqChild = etree.SubElement(seqCTn, qn("p:childTnLst"))

    next_id = 3
    for i, shape in enumerate(shapes_in_order):
        sid = int(shape.shape_id)
        par2 = etree.SubElement(seqChild, qn("p:par"))
        cTn2 = etree.SubElement(par2, qn("p:cTn"), {"id": str(next_id), "fill": "hold"})
        next_id += 1
        stCondLst = etree.SubElement(cTn2, qn("p:stCondLst"))
        etree.SubElement(stCondLst, qn("p:cond"), {"delay": "0" if i == 0 else "indefinite"})
        child2 = etree.SubElement(cTn2, qn("p:childTnLst"))
        par3 = etree.SubElement(child2, qn("p:par"))
        cTn3 = etree.SubElement(par3, qn("p:cTn"), {"id": str(next_id), "fill": "hold"})
        next_id += 1
        st2 = etree.SubElement(cTn3, qn("p:stCondLst"))
        etree.SubElement(st2, qn("p:cond"), {"delay": "0"})
        child3 = etree.SubElement(cTn3, qn("p:childTnLst"))

        animEffect = etree.SubElement(child3, qn("p:animEffect"), {"transition": "in", "filter": "fade"})
        cBhvr = etree.SubElement(animEffect, qn("p:cBhvr"))
        etree.SubElement(cBhvr, qn("p:cTn"), {"id": str(next_id), "dur": "500"})
        next_id += 1
        tgtEl = etree.SubElement(cBhvr, qn("p:tgtEl"))
        etree.SubElement(tgtEl, qn("p:spTgt"), {"spid": str(sid)})

        set_el = etree.SubElement(child3, qn("p:set"))
        cBhvr2 = etree.SubElement(set_el, qn("p:cBhvr"))
        cTnB = etree.SubElement(cBhvr2, qn("p:cTn"), {"id": str(next_id), "dur": "1", "fill": "hold"})
        next_id += 1
        stB = etree.SubElement(cTnB, qn("p:stCondLst"))
        etree.SubElement(stB, qn("p:cond"), {"delay": "0"})
        tgt2 = etree.SubElement(cBhvr2, qn("p:tgtEl"))
        etree.SubElement(tgt2, qn("p:spTgt"), {"spid": str(sid)})
        attr = etree.SubElement(cBhvr2, qn("p:attrNameLst"))
        an = etree.SubElement(attr, qn("p:attrName"))
        an.text = "style.visibility"
        to = etree.SubElement(set_el, qn("p:to"))
        etree.SubElement(to, qn("p:strVal"), {"val": "visible"})


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 Title
    s = prs.slides.add_slide(blank)
    _solid_bg(s, SURFACE)
    add_text(s, Inches(0.7), Inches(0.5), Inches(12), Inches(0.4), [("MANAGER BRIEFING · LEARNWISE IDEAS", 14, True, PRIMARY)])
    add_text(
        s,
        Inches(0.7),
        Inches(1.8),
        Inches(12),
        Inches(2.2),
        [("What we can do", 44, True, PRIMARY), ("with those features", 44, True, INK)],
    )
    add_text(
        s,
        Inches(0.7),
        Inches(4.4),
        Inches(11.5),
        Inches(1.8),
        [
            (
                "We reviewed the LearnWise feature list. We will not copy the whole product. "
                "We take the parts that help students and teachers inside a Moodle class — and build them into EdVidura.",
                18,
                False,
                MUTED,
            ),
            ("October 2026 · Plan: docs/LEARNWISE_IMPLEMENTATION_PLAN.md", 14, False, MUTED),
        ],
    )

    # 2 Two products
    s = prs.slides.add_slide(blank)
    _solid_bg(s, SURFACE)
    add_text(s, Inches(0.7), Inches(0.4), Inches(12), Inches(0.8), [("Two different products", 36, True, INK)])
    left = add_rect(s, Inches(0.8), Inches(1.6), Inches(5.4), Inches(3.4), PAPER, line=MINT)
    add_card_text(
        left,
        "Opens from Moodle. Helps with this class — Ask Vidura, quizzes, study path, learning records.",
        title="EdVidura",
        size=16,
    )
    right = add_rect(s, Inches(7.0), Inches(1.6), Inches(5.4), Inches(3.4), PAPER, line=CORAL)
    add_card_text(
        right,
        "Campus help widget + tutor + grader + ops agent across many systems.",
        title="LearnWise",
        size=16,
        title_color=CORAL,
    )
    add_text(
        s,
        Inches(0.7),
        Inches(5.4),
        Inches(12),
        Inches(1.2),
        [("Same ideas sometimes. Different job. We borrow tutor ideas — we do not rebuild their campus suite.", 18, False, MUTED)],
    )
    add_appear_animation(s, [left, right])

    # 3 Today
    s = prs.slides.add_slide(blank)
    _solid_bg(s, WARM_BG)
    add_text(
        s,
        Inches(0.7),
        Inches(0.35),
        Inches(12),
        Inches(1.0),
        [("What students get today", 34, True, INK), ("Already live or nearly ready in EdVidura.", 18, False, MUTED)],
    )
    cards = [
        (CITRUS, "Ask Vidura", "Answers from this class’s lessons only — with source titles you can check."),
        (MINT, "Smart quizzes", "Different questions per student. Whole book. Difficulty levels. Study-plan check."),
        (SKY, "Voice", "Talk and listen in many Indian languages + English."),
        (CORAL, "Learning trail", "Chat and quiz events saved (xAPI) for teachers and reports."),
    ]
    anim = []
    pos = [(0.7, 1.7), (6.9, 1.7), (0.7, 4.4), (6.9, 4.4)]
    for (x, y), (fill, title, body) in zip(pos, cards):
        sh = add_rect(s, Inches(x), Inches(y), Inches(5.5), Inches(2.2), fill)
        add_card_text(sh, body, title=title, size=15, color=INK, title_color=INK)
        anim.append(sh)
    add_appear_animation(s, anim)

    # 4 Phase A
    s = prs.slides.add_slide(blank)
    _solid_bg(s, DEEP)
    add_text(
        s,
        Inches(0.7),
        Inches(0.35),
        Inches(12),
        Inches(1.0),
        [
            ("What we will add next (Phase A)", 32, True, WHITE),
            ("Small, safe upgrades to Ask Vidura — high value, low risk.", 16, False, MINT),
        ],
    )
    phase_a = [
        ("01  Quick buttons", "Teacher sets 3–5 starter questions so students know what to ask."),
        ("02  Ask again clearly", "If the question is fuzzy, coach asks one short clarifying question first."),
        ("03  Thumbs up / down", "Students rate answers. We save that for teachers and reports."),
        ("04  Flashcards + integrity", "Study cards from lessons. Coach won’t write exams or whole assignments."),
    ]
    anim = []
    for i, (title, body) in enumerate(phase_a):
        x = 0.7 + (i % 2) * 6.2
        y = 1.6 + (i // 2) * 2.5
        sh = add_rect(s, Inches(x), Inches(y), Inches(5.9), Inches(2.2), RGBColor(0x1A, 0x3A, 0x30))
        add_card_text(sh, body, title=title, size=15, color=WHITE, title_color=MINT)
        anim.append(sh)
    add_appear_animation(s, anim)

    # 5 Student day
    s = prs.slides.add_slide(blank)
    _solid_bg(s, SURFACE)
    add_text(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.7), [("A day in the student’s shoes", 34, True, INK)])
    steps = [
        "1. Open Moodle → launch EdVidura for this course",
        "2. Tap a quick button or ask Ask Vidura in their language",
        "3. Get a grounded answer with lesson sources — or a polite “stay on class”",
        "4. Practice — quiz or flashcards on weak spots",
        "5. Thumb the answer — teacher later sees what helped",
    ]
    anim = []
    for i, t in enumerate(steps):
        sh = add_rect(s, Inches(0.8), Inches(1.25 + i * 1.05), Inches(11.7), Inches(0.9), PAPER, line=MINT)
        add_card_text(sh, t, size=17, color=INK)
        anim.append(sh)
    add_appear_animation(s, anim)

    # 6 Teachers
    s = prs.slides.add_slide(blank)
    _solid_bg(s, WARM_BG)
    add_text(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.7), [("What teachers get", 34, True, INK)])
    teacher = [
        (0.7, 1.3, 11.9, 1.2, "Approve the knowledge", "Pick which lessons and manuals the coach may use (SME sources)."),
        (0.7, 2.7, 5.8, 1.7, "Set the coach", "Shortcuts, welcome tone, integrity “don’t do the homework.”"),
        (6.8, 2.7, 5.8, 1.7, "Quiz control", "Level, how hard, how much of the book to cover."),
        (0.7, 4.7, 11.9, 1.7, "Grade help (draft only)", "AI suggests feedback — teacher edits. Nothing auto-sent to Moodle grades."),
    ]
    anim = []
    for x, y, w, h, title, body in teacher:
        sh = add_rect(s, Inches(x), Inches(y), Inches(w), Inches(h), PAPER, line=CITRUS)
        add_card_text(sh, body, title=title, size=15)
        anim.append(sh)
    add_appear_animation(s, anim)

    # 7 Phase B
    s = prs.slides.add_slide(blank)
    _solid_bg(s, SURFACE)
    add_text(
        s,
        Inches(0.7),
        Inches(0.35),
        Inches(12),
        Inches(1.0),
        [("What comes after (Phase B)", 34, True, INK), ("When Phase A works in Riverside.", 18, False, MUTED)],
    )
    phase_b = [
        ("B1  Study plan on screen", "Student sees “what to do next” from adaptive + chat + LMS signals."),
        ("B2  Coach insights", "Teacher view: what students ask, refusals, thumbs, money saved by JEV."),
        ("B3  Better grade assist", "Clearer drafts from rubrics — still copy into Moodle yourself."),
        ("B4  Optional Moodle files", "Pull more course files into the approved knowledge list."),
    ]
    anim = []
    for i, (title, body) in enumerate(phase_b):
        x = 0.7 + (i % 2) * 6.2
        y = 1.6 + (i // 2) * 2.5
        sh = add_rect(s, Inches(x), Inches(y), Inches(5.9), Inches(2.2), PAPER, line=SKY)
        add_card_text(sh, body, title=title, size=15)
        anim.append(sh)
    add_appear_animation(s, anim)

    # 8 Not now
    s = prs.slides.add_slide(blank)
    _solid_bg(s, DEEP)
    add_text(
        s,
        Inches(0.7),
        Inches(0.35),
        Inches(12),
        Inches(1.0),
        [
            ("What we will not build (for now)", 32, True, WHITE),
            ("Say this clearly in meetings — saves months of wrong work.", 16, False, MINT),
        ],
    )
    nos = [
        "Campus helpdesk chatbot / ticket systems",
        "Auto-write grades into Moodle",
        "Ops agent that bulk-changes the LMS",
        "Teams / SharePoint / 400-tool help library",
    ]
    anim = []
    for i, t in enumerate(nos):
        sh = add_rect(s, Inches(1.5), Inches(1.7 + i * 1.2), Inches(10.3), Inches(1.0), CORAL)
        add_card_text(sh, t, size=18, color=WHITE, title_color=WHITE)
        anim.append(sh)
    add_appear_animation(s, anim)

    # 9 Rules
    s = prs.slides.add_slide(blank)
    _solid_bg(s, SURFACE)
    add_text(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.7), [("How we build (safe rules)", 34, True, INK)])
    rules = [
        "1. Reuse modules — coach, SME, quiz, JEV, xAPI — don’t invent a parallel app",
        "2. Decide first, write second — JEV / rules gate; AI writes only when needed",
        "3. One slice at a time — feature flag + test + Riverside demo + sheet update",
        "4. Teacher stays in charge of grades and approved content",
    ]
    anim = []
    for i, t in enumerate(rules):
        sh = add_rect(s, Inches(0.8), Inches(1.4 + i * 1.3), Inches(11.7), Inches(1.1), PAPER, line=PRIMARY)
        add_card_text(sh, t, size=18, color=INK)
        anim.append(sh)
    add_appear_animation(s, anim)

    # 10 Roadmap
    s = prs.slides.add_slide(blank)
    _solid_bg(s, WARM_BG)
    add_text(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.7), [("Roadmap at a glance", 34, True, INK)])
    cols = [
        (MINT, "NOW · A", "Shortcuts, clarify, thumbs, flashcards, integrity, JEV key"),
        (SKY, "NEXT · B", "Study plan UI, coach insights, grade-assist polish"),
        (CITRUS, "LATER · C", "Role-play, video jump-to, nicer extras"),
        (CORAL, "ONLY IF ASKED · D", "Campus widget, help desks, ops writes"),
    ]
    anim = []
    for i, (fill, title, body) in enumerate(cols):
        sh = add_rect(s, Inches(0.55 + i * 3.2), Inches(1.6), Inches(3.0), Inches(4.2), fill)
        add_card_text(sh, body, title=title, size=14, color=INK, title_color=INK)
        anim.append(sh)
    add_appear_animation(s, anim)

    # 11 Ask
    s = prs.slides.add_slide(blank)
    _solid_bg(s, DEEP)
    add_text(s, Inches(0.7), Inches(1.6), Inches(12), Inches(0.5), [("ASK FOR THE ROOM", 14, True, MINT)], align=PP_ALIGN.CENTER)
    add_text(s, Inches(0.7), Inches(2.3), Inches(12), Inches(1.4), [("Approve Phase A?", 44, True, WHITE)], align=PP_ALIGN.CENTER)
    add_text(
        s,
        Inches(1.5),
        Inches(4.0),
        Inches(10.3),
        Inches(1.8),
        [
            (
                "Start with shortcuts, clarify, thumbs, flashcards, and integrity — inside Ask Vidura. Keep Moodle as the gradebook.",
                18,
                False,
                MINT,
            )
        ],
        align=PP_ALIGN.CENTER,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print(f"Wrote {build()}")
